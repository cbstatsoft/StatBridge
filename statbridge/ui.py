# Copyright (C) 2026 Connor Baird
# SPDX-License-Identifier: GPL-3.0-or-later
# This program is free software under GNU GPL version 3 or later.
# It comes WITHOUT ANY WARRANTY; see LICENSE and NOTICE.
from __future__ import annotations
import json
import re
import sys
from os.path import commonprefix
from pathlib import Path
import numpy as np
import pandas as pd
from PySide6.QtCore import (
    Qt,
    QEvent,
    Signal,
    QThread,
    QAbstractTableModel,
    QModelIndex,
    QSettings,
    QStringListModel,
    QSize,
)
from PySide6.QtGui import (
    QAction,
    QActionGroup,
    QKeySequence,
    QFont,
    QFontDatabase,
    QIcon,
    QPainter,
    QColor,
    QSyntaxHighlighter,
    QTextCharFormat,
    QTextCursor,
    QTextBlockFormat,
)
from PySide6.QtWidgets import (
    QApplication,
    QMainWindow,
    QWidget,
    QDockWidget,
    QVBoxLayout,
    QHBoxLayout,
    QLineEdit,
    QPlainTextEdit,
    QTextEdit,
    QListWidget,
    QTableView,
    QTabWidget,
    QLabel,
    QPushButton,
    QFileDialog,
    QMessageBox,
    QToolBar,
    QCompleter,
    QProgressBar,
    QToolButton,
    QCheckBox,
    QHeaderView,
    QDialog,
)
from .engine import Engine
from .catalog import COMMANDS
from .completion import Suggestions
from . import __version__


class CommandLine(QLineEdit):
    def __init__(self, parent=None):
        super().__init__(parent)
        self.setAccessibleName("Stata command")
        self.setAccessibleDescription(
            "Tab completes commands. Shift+Tab leaves this field. F6 moves between workspace panes."
        )
        self.suggestions = Suggestions()
        self.ghost = ""
        self.history_position = 0
        self.draft = ""
        self.completion_edits = {}
        self.setPlaceholderText(
            "Enter a Stata command. Tab: suggestions  •  Right: accept  •  ↑/↓: history"
        )
        self.setFont(QFontDatabase.systemFont(QFontDatabase.SystemFont.FixedFont))
        self.popup_model = QStringListModel(self)
        self.completer = QCompleter(self.popup_model, self)
        self.completer.setCompletionMode(QCompleter.CompletionMode.PopupCompletion)
        self.completer.setWidget(self)
        self.completer.activated.connect(self.accept_completion)
        # QWidget consumes Tab for focus traversal before keyPressEvent.
        # Install after QCompleter, whose popup also filters key events.
        self.installEventFilter(self)
        self.completer.popup().installEventFilter(self)
        self.textChanged.connect(self.refresh)
        self.cursorPositionChanged.connect(self.completion_cursor_changed)

    def refresh(self, text):
        self.completer.popup().hide()
        choices = self.suggestions.candidates(text)
        self.ghost = choices[0][len(text) :] if choices else ""
        self.update()

    def completion_cursor_changed(self, *_):
        self.completer.popup().hide()
        self.update()

    def accept_completion(self, value):
        edit = self.completion_edits.get(value)
        self.setText(value)
        self.setCursorPosition(edit.cursor if edit else len(value))

    def eventFilter(self, watched, event):
        if (
            event.type() == QEvent.Type.KeyPress
            and (
                event.key() == Qt.Key.Key_Tab
                or (
                    self.completer.popup().isVisible()
                    and event.key() in (Qt.Key.Key_Return, Qt.Key.Key_Enter)
                )
            )
            and event.modifiers() == Qt.KeyboardModifier.NoModifier
            and watched in (self, self.completer.popup())
        ):
            self.complete_tab()
            event.accept()
            return True
        return super().eventFilter(watched, event)

    def complete_tab(self):
        popup = self.completer.popup()
        if popup.isVisible():
            index = popup.currentIndex()
            if not index.isValid():
                index = self.completer.completionModel().index(0, 0)
            value = index.data()
            popup.hide()
            if value:
                self.accept_completion(value)
            return
        edits = self.suggestions.complete(self.text(), self.cursorPosition())
        self.completion_edits = {edit.text: edit for edit in edits}
        candidates = list(self.completion_edits)
        if not candidates:
            return
        if len(candidates) == 1:
            self.accept_completion(candidates[0])
            return
        prefix = commonprefix([edit.text[: edit.cursor] for edit in edits])
        if len(prefix) > self.cursorPosition():
            suffix = edits[0].text[edits[0].cursor :]
            self.setText(prefix + suffix)
            self.setCursorPosition(len(prefix))
            return
        self.popup_model.setStringList(candidates)
        self.completer.setCompletionPrefix("")
        self.completer.complete()
        popup.setCurrentIndex(self.completer.completionModel().index(0, 0))

    def paintEvent(self, event):
        super().paintEvent(event)
        if (
            self.ghost
            and self.cursorPosition() == len(self.text())
            and not self.hasSelectedText()
        ):
            fm = self.fontMetrics()
            x = 7 + fm.horizontalAdvance(self.text())
            if x < self.width() - 15:
                painter = QPainter(self)
                painter.setPen(QColor(getattr(self, "ghost_color", "#8794a5")))
                painter.setFont(self.font())
                painter.setClipRect(self.rect().adjusted(4, 2, -5, -2))
                painter.drawText(
                    x, (self.height() + fm.ascent() - fm.descent()) // 2, self.ghost
                )

    def keyPressEvent(self, event):
        key = event.key()
        if (
            key in (Qt.Key.Key_Right, Qt.Key.Key_End)
            and self.cursorPosition() == len(self.text())
            and self.ghost
        ):
            self.setText(self.text() + self.ghost)
            return
        if key == Qt.Key.Key_Tab:
            self.complete_tab()
            return
        if (
            key in (Qt.Key.Key_Up, Qt.Key.Key_Down)
            and not self.completer.popup().isVisible()
        ):
            hist = self.suggestions.history
            if hist:
                if self.history_position == 0:
                    self.draft = self.text()
                self.history_position = max(
                    0,
                    min(
                        len(hist),
                        self.history_position + (1 if key == Qt.Key.Key_Up else -1),
                    ),
                )
                self.setText(
                    hist[-self.history_position]
                    if self.history_position
                    else self.draft
                )
            return
        self.history_position = 0
        super().keyPressEvent(event)


class StataHighlighter(QSyntaxHighlighter):
    def __init__(self, doc):
        super().__init__(doc)
        self.keyword = QTextCharFormat()
        self.keyword.setForeground(QColor("#14679b"))
        self.keyword.setFontWeight(QFont.Weight.Bold)
        self.string = QTextCharFormat()
        self.string.setForeground(QColor("#aa4c2a"))
        self.comment = QTextCharFormat()
        self.comment.setForeground(QColor("#678069"))
        self.macro = QTextCharFormat()
        self.macro.setForeground(QColor("#9059a3"))

    def highlightBlock(self, text):
        for m in re.finditer(
            r"\b(?:"
            + "|".join(COMMANDS)
            + r"|if|in|using|quietly|capture|by|bysort)\b",
            text,
        ):
            self.setFormat(m.start(), m.end() - m.start(), self.keyword)
        for m in re.finditer(r'"[^"]*"', text):
            self.setFormat(m.start(), m.end() - m.start(), self.string)
        for m in re.finditer(r"`\w+'|\$\w+", text):
            self.setFormat(m.start(), m.end() - m.start(), self.macro)
        if text.lstrip().startswith("*"):
            self.setFormat(0, len(text), self.comment)
        else:
            quote = False
            for i, c in enumerate(text):
                if c == '"':
                    quote = not quote
                if not quote and text[i : i + 2] == "//":
                    self.setFormat(i, len(text) - i, self.comment)
                    break


class FrameModel(QAbstractTableModel):
    editRequested = Signal(str)

    def __init__(self, frame=None, editable=False):
        super().__init__()
        self.frame = frame if frame is not None else pd.DataFrame()
        self.editable = editable
        self.formats = {}

    def replace(self, frame, formats=None):
        self.formats = dict(formats or {})
        self.beginResetModel()
        self.frame = frame
        self.endResetModel()

    def rowCount(self, parent=QModelIndex()):
        return 0 if parent.isValid() else len(self.frame)

    def columnCount(self, parent=QModelIndex()):
        return 0 if parent.isValid() else len(self.frame.columns)

    def data(self, index, role=Qt.ItemDataRole.DisplayRole):
        if not index.isValid():
            return None
        value = self.frame.iat[index.row(), index.column()]
        if role in (Qt.ItemDataRole.DisplayRole, Qt.ItemDataRole.EditRole):
            if pd.isna(value):
                return "."
            if role == Qt.ItemDataRole.DisplayRole:
                from .dates import display_date
                formatted = display_date(value, self.formats.get(self.frame.columns[index.column()]))
                if formatted is not None:
                    return formatted
            if isinstance(value, (float, np.floating)):
                return format(value, ".10g")
            return str(value)
        if role == Qt.ItemDataRole.TextAlignmentRole and isinstance(
            value, (int, float, np.number)
        ):
            return int(Qt.AlignmentFlag.AlignRight | Qt.AlignmentFlag.AlignVCenter)
        if role == Qt.ItemDataRole.ForegroundRole and pd.isna(value):
            return QColor("#99a1ad")

    def headerData(self, section, orientation, role=Qt.ItemDataRole.DisplayRole):
        if role == Qt.ItemDataRole.DisplayRole:
            return (
                str(self.frame.columns[section])
                if orientation == Qt.Orientation.Horizontal
                else str(section + 1)
            )

    def flags(self, index):
        flags = super().flags(index)
        return flags | Qt.ItemFlag.ItemIsEditable if self.editable else flags

    def setData(self, index, value, role=Qt.ItemDataRole.EditRole):
        if role != Qt.ItemDataRole.EditRole or not self.editable:
            return False
        col = self.frame.columns[index.column()]
        if pd.api.types.is_numeric_dtype(self.frame[col]):
            try:
                text = "." if str(value) in ("", ".") else str(float(value))
            except ValueError:
                return False
        else:
            if '"' in str(value) or "\n" in str(value):
                return False
            text = '"' + str(value) + '"'
        self.editRequested.emit(f"replace {col} = {text} in {index.row() + 1}")
        return False  # Engine transaction updates the displayed model when complete.


class Worker(QThread):
    event = Signal(object)

    def __init__(self, engine, source):
        super().__init__()
        self.engine = engine
        self.source = source

    def run(self):
        self.engine.emit_callback = self.event.emit
        try:
            self.engine.run(self.source)
        except Exception as exc:
            self.event.emit({"kind": "error", "text": str(exc), "payload": None})
        finally:
            self.engine.emit_callback = None


class MainWindow(QMainWindow):
    def __init__(self, syntax=None):
        super().__init__()
        self.settings = QSettings("StatBridge", "StatBridge Qt")
        self.dark_mode = self.settings.value("dark_mode", False, type=bool)
        self.remember_history = self.settings.value("remember_history", True, type=bool)
        self.remember_cwd = self.settings.value("remember_cwd", True, type=bool)
        chosen_syntax = syntax or self.settings.value("syntax_mode", "stata")
        if chosen_syntax not in ("stata", "logos"):
            chosen_syntax = "stata"
        self.engine = Engine(syntax_mode=chosen_syntax)
        self.worker = None
        self.pending_graph = None
        self.do_path = None
        if self.remember_cwd:
            cwd = self.settings.value("cwd", "")
            if cwd and Path(cwd).is_dir():
                self.engine.cwd = Path(cwd)
        else:
            self.settings.remove("cwd")
            self.engine.cwd = Path.home()
        self.setWindowTitle(f"StatBridge Qt {__version__}")
        self.resize(1400, 900)
        self.setDockOptions(
            QMainWindow.DockOption.AllowNestedDocks
            | QMainWindow.DockOption.AllowTabbedDocks
        )
        fixed = QFontDatabase.systemFont(QFontDatabase.SystemFont.FixedFont)
        fixed.setPointSize(11)
        self.tabs = QTabWidget()
        self.setCentralWidget(self.tabs)
        self.results = QTextEdit()
        self.results_plain = []
        self.results_records = []
        self.results.setReadOnly(True)
        self.results.setFont(fixed)
        self.results.setLineWrapMode(QTextEdit.LineWrapMode.NoWrap)
        self.tabs.addTab(self.results, "Results")
        self.data_view = QTableView()
        self.data_model = FrameModel(editable=True)
        self.data_view.setModel(self.data_model)
        self.data_view.setAlternatingRowColors(True)
        self.data_view.horizontalHeader().setSectionResizeMode(
            QHeaderView.ResizeMode.Interactive
        )
        self.tabs.addTab(self.data_view, "Data editor")
        self.data_model.editRequested.connect(self.start)
        editor_page = QWidget()
        el = QVBoxLayout(editor_page)
        el.setContentsMargins(8, 8, 8, 8)
        actions = QHBoxLayout()
        self.editor_name = QLabel("Untitled.do")
        actions.addWidget(self.editor_name)
        actions.addStretch()
        self.editor_buttons = []
        for _ in range(3):
            b = QToolButton()
            b.setToolButtonStyle(Qt.ToolButtonStyle.ToolButtonTextBesideIcon)
            b.setFocusPolicy(Qt.FocusPolicy.StrongFocus)
            self.editor_buttons.append(b)
            actions.addWidget(b)
        el.addLayout(actions)
        self.editor = QPlainTextEdit()
        self.editor.setFont(fixed)
        self.editor.setLineWrapMode(QPlainTextEdit.LineWrapMode.NoWrap)
        self.editor.setTabStopDistance(
            4 * self.editor.fontMetrics().horizontalAdvance(" ")
        )
        self.highlighter = StataHighlighter(self.editor.document())
        el.addWidget(self.editor)
        self.editor.cursorPositionChanged.connect(self.editor_location)
        self.editor.document().modificationChanged.connect(self.editor_title)
        self.tabs.addTab(editor_page, "Do-file editor")
        self.graph_page = QWidget()
        self.graph_layout = QVBoxLayout(self.graph_page)
        self.graph_layout.addWidget(
            QLabel("Create a graph with histogram, scatter or line.")
        )
        self.tabs.addTab(self.graph_page, "Graphs")
        self.audit = QTableView()
        self.audit_model = FrameModel()
        self.audit.setModel(self.audit_model)
        self.audit.setWordWrap(True)
        self.audit.horizontalHeader().setStretchLastSection(True)
        self.tabs.addTab(self.audit, "Compatibility")
        self.review = QListWidget()
        self.review.itemDoubleClicked.connect(
            lambda item: self.command.setText(item.text())
        )
        review_panel = QWidget()
        review_layout = QVBoxLayout(review_panel)
        review_layout.setContentsMargins(4, 4, 4, 4)
        review_layout.addWidget(self.review)
        self.clear_history_button = QToolButton()
        self.clear_history_button.setToolButtonStyle(
            Qt.ToolButtonStyle.ToolButtonTextOnly
        )
        review_layout.addWidget(self.clear_history_button)
        self.remember_history_box = QCheckBox("Remember history")
        self.remember_history_box.setToolTip(
            "Save command history between sessions. Turning this off removes previously saved history."
        )
        review_layout.addWidget(self.remember_history_box)
        self.dock("Review", review_panel, Qt.DockWidgetArea.LeftDockWidgetArea)
        right = QWidget()
        rl = QVBoxLayout(right)
        rl.setContentsMargins(5, 5, 5, 5)
        self.var_search = QLineEdit()
        self.var_search.setPlaceholderText("Filter variables")
        self.var_search.setAccessibleName("Filter variables")
        rl.addWidget(self.var_search)
        self.variables = QListWidget()
        rl.addWidget(self.variables)
        self.var_details = QLabel("No dataset loaded")
        self.var_details.setWordWrap(True)
        rl.addWidget(self.var_details)
        self.var_search.textChanged.connect(self.filter_variables)
        self.variables.itemClicked.connect(self.variable_info)
        self.variables.itemDoubleClicked.connect(self.insert_variable)
        self.dock(
            "Variables / Properties", right, Qt.DockWidgetArea.RightDockWidgetArea
        )
        cmd_widget = QWidget()
        cl = QVBoxLayout(cmd_widget)
        cl.setContentsMargins(8, 6, 8, 6)
        row = QHBoxLayout()
        prompt = QLabel(".")
        prompt.setFont(fixed)
        row.addWidget(prompt)
        self.command = CommandLine()
        self.command.setMinimumHeight(35)
        row.addWidget(self.command)
        self.run_button = QPushButton("Run")
        self.run_button.setProperty("primary", True)
        self.run_button.setToolTip("Execute the command in the Command field (Enter)")
        self.run_button.clicked.connect(self.submit)
        row.addWidget(self.run_button)
        self.stop_button = QPushButton("Stop")
        self.stop_button.setToolTip(
            "Stop after the current backend call (Ctrl+Break or Ctrl+Shift+.)"
        )
        self.stop_button.clicked.connect(self.stop)
        self.stop_button.setEnabled(False)
        row.addWidget(self.stop_button)
        cl.addLayout(row)
        self.command.returnPressed.connect(self.submit)
        self.command.textChanged.connect(self.command_help)
        self.hint = QLabel("Ready. Type help to list supported commands.")
        cl.addWidget(self.hint)
        self.dock("Command", cmd_widget, Qt.DockWidgetArea.BottomDockWidgetArea)
        self.progress = QProgressBar()
        self.progress.setMaximumWidth(160)
        self.progress.hide()
        self.statusBar().addPermanentWidget(self.progress)
        from .learning_ui import LearningPanel

        self.learning = LearningPanel(self.engine)
        self.learning.openExample.connect(self.open_learning_example)
        self.learning_dock = self.dock(
            "Learn", self.learning, Qt.DockWidgetArea.RightDockWidgetArea
        )
        self.learning_dock.visibilityChanged.connect(self.size_learning_dock)
        self.learning_dock.hide()
        self.make_menus()
        self.apply_style()
        if self.remember_history:
            raw = self.settings.value("history", "[]")
            try:
                saved = json.loads(raw)
                if isinstance(saved, list):
                    self.command.suggestions.history = [
                        cmd for cmd in saved if isinstance(cmd, str)
                    ][-1000:]
            except (ValueError, TypeError):
                pass
        else:
            self.settings.remove("history")
        for cmd in self.command.suggestions.history:
            self.review.addItem(cmd)
        geom = self.settings.value("geometry")
        state = self.settings.value("layout")
        if geom:
            self.restoreGeometry(geom)
        if state:
            self.restoreState(state)
        self.learning_dock.setVisible(
            self.settings.value("learning_enabled", False, type=bool)
        )
        self.append_output(
            "StatBridge Qt  "
            + __version__
            + "\nStata syntax • Python statistics • Explicit compatibility reporting\n\nType help for commands, or open examples/demo.do in the Do-file editor.\nOptional lessons and explanations: click Learn in the toolbar.\nAll estimation output comes from Python backends. Unsupported syntax stops execution.\n",
            "output",
        )
        from .legal import NOTICE
        self.append_output("Public beta. " + NOTICE + "\n", "output")
        self.refresh()
        self.command.setFocus()

    def show_text_dialog(self, title, text):
        from PySide6.QtWidgets import QDialog, QDialogButtonBox
        dialog = QDialog(self)
        dialog.setWindowTitle(title)
        dialog.resize(760, 540)
        layout = QVBoxLayout(dialog)
        viewer = QPlainTextEdit()
        viewer.setReadOnly(True)
        viewer.setPlainText(text)
        layout.addWidget(viewer)
        buttons = QDialogButtonBox(QDialogButtonBox.StandardButton.Close)
        copy_button = buttons.addButton("Copy", QDialogButtonBox.ButtonRole.ActionRole)
        copy_button.clicked.connect(lambda: QApplication.clipboard().setText(text))
        buttons.rejected.connect(dialog.reject)
        layout.addWidget(buttons)
        dialog.exec()

    def show_licence(self):
        from .legal import license_text
        self.show_text_dialog("Licence and third-party notices", license_text())

    def show_diagnostics(self):
        from .diagnostics import report, format_report
        self.show_text_dialog("Installation diagnostics", format_report(report()))

    def dock(self, title, widget, area):
        dock = QDockWidget(title, self)
        dock.setObjectName(title)
        dock.setWidget(widget)
        self.addDockWidget(area, dock)
        return dock

    def make_menus(self):
        file = self.menuBar().addMenu("&File")
        data = self.menuBar().addMenu("&Data")
        analysis = self.menuBar().addMenu("&Statistics")
        view = self.menuBar().addMenu("&Window")
        helpmenu = self.menuBar().addMenu("&Help")
        self.actions = []

        def action(menu, label, slot, shortcut=None):
            a = QAction(label, self)
            a.triggered.connect(slot)
            if shortcut:
                a.setShortcut(shortcut)
            menu.addAction(a)
            self.actions.append(a)
            return a

        action(helpmenu,"Workflow coverage and limits",lambda: self.start("coverage"))
        open_action = action(
            file, "Open dataset…", self.open_data, QKeySequence.StandardKey.Open
        )
        save_action = action(
            file, "Save dataset as…", self.save_data, QKeySequence.StandardKey.SaveAs
        )
        action(file, "Import CSV…", self.import_csv)
        action(file, "Import Excel…", self.import_excel)
        action(file, "Import World Bank panel…", self.import_databank)
        action(file, "Export dataset formats…", self.export_statistical)
        action(file, "Export Excel…", self.export_excel)
        results_menu = self.menuBar().addMenu("&Results")
        self.copy_table_action = action(results_menu, "Copy latest table", self.copy_latest_table)
        self.export_table_action = action(results_menu, "Export latest table…", self.export_latest_table)
        self.folder_action = action(
            file, "Choose working folder…", self.choose_working_folder
        )
        action(file, "Use do-file folder", self.use_do_folder)
        action(file, "List working folder", lambda: self.start("ls"))
        self.remember_cwd_action = QAction(
            "Remember working folder between sessions", self
        )
        self.remember_cwd_action.setCheckable(True)
        self.remember_cwd_action.setChecked(self.remember_cwd)
        self.remember_cwd_action.setToolTip(
            "When disabled, new sessions start in your home folder. The current folder stays unchanged."
        )
        self.remember_cwd_action.toggled.connect(self.set_remember_cwd)
        file.addAction(self.remember_cwd_action)
        file.addSeparator()
        open_do_action = action(file, "Open do-file…", self.open_do, "Ctrl+Shift+O")
        save_do_action = action(
            file, "Save do-file", self.save_do, QKeySequence.StandardKey.Save
        )
        action(file, "Save results…", self.save_results)
        action(file, "Export compatibility report…", self.export_audit)
        file.addSeparator()
        action(file, "Exit", self.close, "Ctrl+Q")
        action(data, "Describe", lambda: self.start("describe"))
        action(data, "Summarize", lambda: self.start("summarize"))
        action(data, "Data editor", lambda: self.tabs.setCurrentIndex(1))
        from .analysis_ui import AnalysisDialog

        self.analysis_actions = {}
        for kind, title in AnalysisDialog.TITLES.items():
            self.analysis_actions[kind] = action(
                analysis,
                title + "…",
                lambda checked=False, k=kind: self.open_analysis(k),
            )
        for d in self.findChildren(QDockWidget):
            view.addAction(d.toggleViewAction())
        view.addSeparator()
        history_menu = view.addMenu("Command history")
        self.clear_history_action = action(
            history_menu, "Clear history", self.clear_history
        )
        self.clear_history_action.setToolTip(
            "Clear the Review pane, command recall and saved history"
        )
        self.clear_history_button.setDefaultAction(self.clear_history_action)
        self.remember_history_action = QAction(
            "Remember history between sessions", self
        )
        self.remember_history_action.setCheckable(True)
        self.remember_history_action.setChecked(self.remember_history)
        self.remember_history_action.toggled.connect(self.set_remember_history)
        history_menu.addAction(self.remember_history_action)
        self.remember_history_box.setChecked(self.remember_history)
        self.remember_history_box.toggled.connect(
            self.remember_history_action.setChecked
        )
        self.remember_history_action.toggled.connect(
            self.remember_history_box.setChecked
        )
        self.review.setContextMenuPolicy(Qt.ContextMenuPolicy.ActionsContextMenu)
        self.review.addAction(self.clear_history_action)
        self.review.addAction(self.remember_history_action)
        language_menu = view.addMenu("Command language")
        self.language_group = QActionGroup(self)
        self.language_group.setExclusive(True)
        self.stata_language_action = QAction("Stata-compatible syntax", self)
        self.logos_language_action = QAction("Logos — English / formal logic", self)
        for mode, item in (("stata", self.stata_language_action), ("logos", self.logos_language_action)):
            item.setCheckable(True)
            item.setChecked(self.engine.syntax_mode == mode)
            item.triggered.connect(lambda checked=False, m=mode: self.set_command_language(m))
            self.language_group.addAction(item)
            language_menu.addAction(item)
            self.actions.append(item)
        action(helpmenu, "Logos language guide", lambda: self.start("stata: help logos"))
        self.dark_action = QAction("Material dark mode", self)
        self.dark_action.setCheckable(True)
        self.dark_action.setChecked(self.dark_mode)
        self.dark_action.toggled.connect(self.set_dark_mode)
        view.addAction(self.dark_action)
        action(helpmenu, "Command reference", lambda: self.start("help"), "F1")
        action(helpmenu, "Compatibility notes", lambda: self.start("compatibility"))
        action(helpmenu, "About / backend versions", lambda: self.start("about"))
        action(helpmenu, "Licence and third-party notices", self.show_licence)
        action(helpmenu, "Installation diagnostics", self.show_diagnostics)
        toolbar = QToolBar("Main")
        self.toolbar = toolbar
        toolbar.setObjectName("Main")
        toolbar.setToolButtonStyle(Qt.ToolButtonStyle.ToolButtonTextBesideIcon)
        toolbar.setIconSize(QSize(20, 20))
        self.addToolBar(toolbar)
        self.icon_controls = [
            (open_action, "open"),
            (save_action, "save"),
            (self.folder_action, "folder"),
            (open_do_action, "script"),
            (save_do_action, "save"),
            (self.run_button, "run"),
            (self.stop_button, "stop"),
            (self.clear_history_action, "clear"),
        ]
        for a, title in (
            (open_action, "Open data"),
            (save_action, "Save data"),
            (self.folder_action, "Folder"),
        ):
            a.setIconText(title)
            toolbar.addAction(a)
        toolbar.addSeparator()
        self.analysis_button = QToolButton()
        self.analysis_button.setText("Analyse")
        self.analysis_button.setAccessibleName("Choose analysis")
        self.analysis_button.setToolTip("Choose an analysis and select variables")
        self.analysis_button.setToolButtonStyle(
            Qt.ToolButtonStyle.ToolButtonTextBesideIcon
        )
        self.analysis_button.setPopupMode(QToolButton.ToolButtonPopupMode.InstantPopup)
        self.analysis_button.setMenu(analysis)
        toolbar.addWidget(self.analysis_button)
        self.icon_controls.append((self.analysis_button, "analyse"))
        self.run_editor_action = action(
            file, "Run do-file selection / all", self.run_editor, "Ctrl+D"
        )
        self.run_editor_action.setIconText("Run do-file")
        self.run_editor_action.setToolTip(
            "Run selected do-file text, or the whole do-file when nothing is selected (Ctrl+D)"
        )
        toolbar.addAction(self.run_editor_action)
        self.stop_action = QAction("Stop", self)
        self.stop_action.setToolTip(
            "Stop after the current backend call (Ctrl+Break or Ctrl+Shift+.)"
        )
        self.stop_action.setShortcuts(
            [QKeySequence("Ctrl+Break"), QKeySequence("Ctrl+Shift+.")]
        )
        self.stop_action.triggered.connect(self.stop)
        self.stop_action.setEnabled(False)
        toolbar.addAction(self.stop_action)
        file.addAction(self.stop_action)
        self.icon_controls.extend(
            [(self.run_editor_action, "run"), (self.stop_action, "stop")]
        )
        toolbar.addSeparator()
        self.learning_action = self.learning_dock.toggleViewAction()
        self.learning_action.setText("Learn")
        self.learning_action.setToolTip(
            "Show or hide optional lessons and command explanations"
        )
        toolbar.addAction(self.learning_action)
        helpmenu.addAction(self.learning_action)
        self.icon_controls.append((self.learning_action, "learn"))
        self.theme_button = QToolButton()
        self.theme_button.setDefaultAction(self.dark_action)
        self.theme_button.setToolButtonStyle(Qt.ToolButtonStyle.ToolButtonIconOnly)
        self.theme_button.setFixedSize(36, 36)
        toolbar.addWidget(self.theme_button)
        for b, a, title in zip(
            self.editor_buttons,
            (open_do_action, save_do_action, self.run_editor_action),
            ("Open", "Save", "Run do-file"),
        ):
            b.setDefaultAction(a)
            b.setText(title)
        for b in self.findChildren(QToolButton):
            b.setFocusPolicy(Qt.FocusPolicy.StrongFocus)
            b.setMinimumHeight(32)
        action(view, "Focus command", lambda: self.command.setFocus(), "Ctrl+L")
        action(view, "Next workspace pane", self.focus_next_pane, "F6")
        self.results.setAccessibleName("Analysis results")
        self.editor.setAccessibleName("Do-file editor")
        self.review.setAccessibleName("Command history")
        self.variables.setAccessibleName("Dataset variables")
        self.data_view.setAccessibleName("Dataset editor")

    def focus_next_pane(self):
        current = self.tabs.currentWidget()
        target = self.editor if self.tabs.currentIndex() == 2 else current
        panes = [self.command, target, self.review, self.var_search]
        focus = QApplication.focusWidget()
        index = (
            next(
                (i for i, p in enumerate(panes) if p is focus or p.isAncestorOf(focus)),
                -1,
            )
            if focus
            else -1
        )
        panes[(index + 1) % len(panes)].setFocus()

    def open_analysis(self, kind):
        if self.worker:
            return
        from .analysis_ui import AnalysisDialog

        dialog = AnalysisDialog(self.engine, kind, self)
        if dialog.exec() == QDialog.DialogCode.Accepted:
            if dialog.insert_only:
                self.editor.appendPlainText(dialog.source)
                self.tabs.setCurrentIndex(2)
                self.editor.setFocus()
            else:
                self.tabs.setCurrentIndex(0)
                self.start(dialog.source)

    def save_working_folder(self):
        if self.remember_cwd:
            self.settings.setValue("cwd", str(self.engine.cwd))
        else:
            self.settings.remove("cwd")
        self.settings.sync()

    def set_remember_cwd(self, enabled):
        self.remember_cwd = bool(enabled)
        self.settings.setValue("remember_cwd", self.remember_cwd)
        self.save_working_folder()

    def save_history(self):
        if self.remember_history:
            self.settings.setValue(
                "history", json.dumps(self.command.suggestions.history[-1000:])
            )
        else:
            self.settings.remove("history")
        self.settings.sync()

    def set_remember_history(self, enabled):
        self.remember_history = bool(enabled)
        self.settings.setValue("remember_history", self.remember_history)
        self.save_history()

    def clear_history(self):
        self.command.suggestions.history.clear()
        # Keep the engine audit intact; its previous commands must not reappear on refresh.
        self.history_offset = len(self.engine.history)
        self.review.clear()
        self.command.history_position = 0
        self.command.draft = self.command.text()
        self.command.completer.popup().hide()
        self.command.refresh(self.command.text())
        self.settings.remove("history")
        self.settings.sync()
        self.statusBar().showMessage("Command history cleared.", 4000)

    def set_command_language(self, mode):
        if self.worker is not None or mode not in ("stata", "logos"):
            return
        self.engine.syntax_mode = mode
        self.settings.setValue("syntax_mode", mode)
        self.statusBar().showMessage(f"Command language: {mode}")
        self.refresh()

    def set_dark_mode(self, enabled):
        self.dark_mode = bool(enabled)
        self.settings.setValue("dark_mode", self.dark_mode)
        self.apply_style()
        scrollbar = self.results.verticalScrollBar()
        position = scrollbar.value()
        at_end = position == scrollbar.maximum()
        self.results.clear()
        for text, kind, html in self.results_records:
            self.append_output(text, kind, html, record=False)
        scrollbar.setValue(scrollbar.maximum() if at_end else position)

    def apply_style(self):
        from .themes import DARK, LIGHT, palette, stylesheet, theme_icon

        self.colors = DARK if self.dark_mode else LIGHT
        self.dark_action.setIcon(theme_icon(self.dark_mode, self.colors["accent"]))
        theme_label = (
            "Switch to light mode" if self.dark_mode else "Switch to dark mode"
        )
        self.dark_action.setToolTip(theme_label)
        self.theme_button.setAccessibleName(theme_label)
        from .icons import action_icon

        for control, name in self.icon_controls:
            control.setIcon(
                action_icon(name, self.colors["text"], self.colors["muted"])
            )
        self.setPalette(palette(self.colors))
        self.setStyleSheet(stylesheet(self.colors))
        popup = self.command.completer.popup()
        popup.setPalette(palette(self.colors))
        popup.setStyleSheet(stylesheet(self.colors))
        self.command.ghost_color = self.colors["ghost"]
        self.command.update()
        for name in ("keyword", "string", "comment", "macro"):
            color = self.colors["accent" if name == "keyword" else name]
            getattr(self.highlighter, name).setForeground(QColor(color))
        self.highlighter.rehighlight()
        self.learning.apply_theme(self.colors)

    def append_output(self, text, kind, html=None, record=True):
        if record:
            self.results_plain.append(text + "\n")
            self.results_records.append((text, kind, html))
        if kind in ("warning", "error"):
            import textwrap

            width = max(
                45,
                int(
                    self.results.viewport().width()
                    / self.results.fontMetrics().horizontalAdvance("M")
                )
                - 3,
            )
            text = "\n".join(
                textwrap.fill(line, width=width, subsequent_indent="  ")
                for line in text.splitlines()
            )
        cursor = self.results.textCursor()
        cursor.movePosition(QTextCursor.MoveOperation.End)
        fmt = QTextCharFormat()
        fmt.setFont(self.results.font())
        fmt.setForeground(QColor(self.colors.get(kind, self.colors["text"])))
        fmt.setFontWeight(
            QFont.Weight.Bold if kind == "command" else QFont.Weight.Normal
        )
        if html:
            from .themes import report_html

            cursor.setCharFormat(fmt)
            cursor.insertHtml(report_html(html, self.colors))
            cursor.movePosition(QTextCursor.MoveOperation.End)
            cursor.insertBlock(QTextBlockFormat(), fmt)
        else:
            cursor.insertText(text + "\n", fmt)
        self.results.setTextCursor(cursor)
        self.results.ensureCursorVisible()

    def event_received(self, event):
        kind = event["kind"]
        if kind in ("command", "error", "output", "warning", "translation"):
            payload = event.get("payload")
            html = (
                payload.get("html")
                if kind == "output" and isinstance(payload, dict)
                else None
            )
            self.append_output(event["text"], kind, html)
        elif kind == "graph":
            self.pending_graph = event["payload"]
        elif kind == "progress":
            self.statusBar().showMessage(event["text"])

    def submit(self):
        text = self.command.text().strip()
        if text:
            self.command.clear()
            self.start(text)

    def start(self, source):
        if self.worker is not None:
            return
        self.learning.set_busy(True)
        self.engine.stop.clear()
        self.pending_graph = None
        self.worker = Worker(self.engine, source)
        self.worker.event.connect(self.event_received)
        self.worker.finished.connect(self.finished)
        self.command.setEnabled(False)
        self.run_button.setEnabled(False)
        self.stop_button.setEnabled(True)
        self.stop_action.setEnabled(True)
        self.analysis_button.setEnabled(False)
        self.data_view.setEnabled(False)
        for a in self.actions:
            a.setEnabled(False)
        self.progress.setRange(0, 0)
        self.progress.show()
        self.statusBar().showMessage(
            "Running. Stop takes effect between commands or bootstrap replications."
        )
        self.worker.start()

    def finished(self):
        mode = self.engine.syntax_mode
        self.settings.setValue("syntax_mode", mode)
        for name, item in (("stata", self.stata_language_action), ("logos", self.logos_language_action)):
            item.blockSignals(True)
            item.setChecked(name == mode)
            item.blockSignals(False)
        self.worker.deleteLater()
        self.worker = None
        self.learning.set_busy(False)
        self.command.setEnabled(True)
        self.run_button.setEnabled(True)
        self.stop_button.setEnabled(False)
        self.stop_action.setEnabled(False)
        self.analysis_button.setEnabled(True)
        self.data_view.setEnabled(True)
        for a in self.actions:
            a.setEnabled(True)
        self.progress.hide()
        self.refresh()
        if self.pending_graph is not None:
            self.show_graph(self.pending_graph)
            self.pending_graph = None
        self.command.setFocus()

    def stop(self):
        if self.worker:
            self.engine.stop.set()
            self.statusBar().showMessage(
                "Stop requested. Waiting for current backend call to finish…"
            )

    def size_learning_dock(self, visible):
        if visible:
            self.resizeDocks(
                [self.learning_dock],
                [int(self.height() * 0.65)],
                Qt.Orientation.Vertical,
            )

    def refresh(self):
        has_table = getattr(self.engine, "results_table", None) is not None and self.worker is None
        self.copy_table_action.setEnabled(has_table)
        self.export_table_action.setEnabled(has_table)
        self.learning.refresh()
        self.data_model.replace(self.engine.df.copy(deep=False), self.engine.formats)
        self.variables.clear()
        self.variables.addItems(list(self.engine.df.columns))
        self.var_details.setText(
            f"{len(self.engine.df):,} observations, {len(self.engine.df.columns)} variables.\n"
            "Select a variable for properties."
        )
        self.filter_variables(self.var_search.text())
        hist = self.command.suggestions.history
        # Only append new engine commands. Persist a bounded history, not datasets.
        offset = getattr(self, "history_offset", 0)
        hist.extend(self.engine.history[offset:])
        self.history_offset = len(self.engine.history)
        hist[:] = hist[-1000:]
        self.review.clear()
        self.review.addItems(hist)
        self.review.scrollToBottom()
        self.command.suggestions.sync(self.engine)
        self.command.refresh(self.command.text())
        self.audit_model.replace(
            pd.DataFrame(self.engine.notices, columns=["code", "command", "message"])
        )
        self.audit.setColumnWidth(0, 130)
        self.audit.setColumnWidth(1, 260)
        self.audit.resizeRowsToContents()
        self.statusBar().showMessage(
            f"{len(self.engine.df):,} observations  |  {len(self.engine.df.columns)} variables  |  {self.engine.cwd}"
            + ("  |  Unsaved data changes" if self.engine.dirty else "")
        )
        self.setWindowTitle(
            f"{self.engine.filename.name if self.engine.filename else 'Untitled'}{' *' if self.engine.dirty else ''} - StatBridge Qt {__version__}"
        )

    def show_graph(self, fig):
        from matplotlib.backends.backend_qtagg import (
            FigureCanvasQTAgg,
            NavigationToolbar2QT,
        )

        while self.graph_layout.count():
            item = self.graph_layout.takeAt(0)
            if item.widget():
                item.widget().deleteLater()
        canvas = FigureCanvasQTAgg(fig)
        self.graph_layout.addWidget(NavigationToolbar2QT(canvas, self))
        self.graph_layout.addWidget(canvas)
        canvas.draw()
        self.tabs.setCurrentIndex(3)

    def filter_variables(self, text):
        for i in range(self.variables.count()):
            item = self.variables.item(i)
            item.setHidden(text.lower() not in item.text().lower())

    def variable_info(self, item):
        c = item.text()
        s = self.engine.df[c]
        self.var_details.setText(
            f"{c}\n{self.engine.labels.get(c, '')}\nType: {s.dtype}\nMissing: {s.isna().sum():,}\nUnique: {s.nunique():,}"
        )

    def insert_variable(self, item):
        if not self.worker:
            self.command.insert(
                (
                    " "
                    if self.command.text() and not self.command.text().endswith(" ")
                    else ""
                )
                + item.text()
            )
            self.command.setFocus()

    def command_help(self, text):
        from .catalog import ALIASES

        key = text.split(" ", 1)[0].split(",")[0]
        key = ALIASES.get(key, key)
        self.hint.setText(
            COMMANDS.get(
                key,
                (
                    "Tab: completion list. Right arrow: accept inline suggestion. Up/down: history.",
                ),
            )[0]
        )

    def prefill(self, text):
        self.command.setText(text)
        self.command.setFocus()

    def quote(self, path):
        if '"' in path:
            raise ValueError("Paths containing double quotes are unsupported")
        return '"' + path + '"'

    def discard_data(self):
        return (
            not self.engine.dirty
            or QMessageBox.question(
                self,
                "Replace unsaved dataset?",
                "The in-memory dataset has unsaved changes. Discard those changes?",
                QMessageBox.StandardButton.Discard | QMessageBox.StandardButton.Cancel,
                QMessageBox.StandardButton.Cancel,
            )
            == QMessageBox.StandardButton.Discard
        )

    def import_databank(self):
        if self.worker: return
        from .databank_ui import DatabankDialog
        dialog = DatabankDialog(cwd=self.engine.cwd, parent=self)
        if dialog.exec() == QDialog.DialogCode.Accepted and self.discard_data():
            self.start(dialog.command())

    def export_statistical(self):
        if self.worker: return
        from .format_dialog import FormatDialog
        dialog = FormatDialog(self.engine.cwd, self.engine.df.columns, exporting=True, parent=self)
        if dialog.exec() == QDialog.DialogCode.Accepted:
            self.start(dialog.command())

    def import_excel(self):
        if self.worker: return
        path, _ = QFileDialog.getOpenFileName(self, "Import Excel", str(self.engine.cwd), "Excel (*.xlsx)")
        if not path: return
        try:
            from openpyxl import load_workbook
            from .io_dialogs import ExcelImportDialog
            book = load_workbook(path, read_only=True, data_only=True)
            sheets = book.sheetnames
            book.close()
            dialog = ExcelImportDialog(path, sheets, self)
            if dialog.exec() == QDialog.DialogCode.Accepted and self.discard_data():
                self.start(dialog.command())
        except Exception as error:
            QMessageBox.warning(self, "Excel import", str(error))

    def export_excel(self):
        if self.worker: return
        path, _ = QFileDialog.getSaveFileName(self, "Export Excel", str(self.engine.cwd / "data.xlsx"), "Excel (*.xlsx)")
        if path:
            self.start("export excel using " + self.quote(path) + ", firstrow(variables) replace")

    def copy_latest_table(self):
        table = getattr(self.engine, "results_table", None)
        if self.worker or table is None: return
        QApplication.clipboard().setText(table.to_csv(sep="\t"))
        self.statusBar().showMessage("Latest results table copied", 4000)

    def export_latest_table(self):
        if self.worker or getattr(self.engine, "results_table", None) is None: return
        path, _ = QFileDialog.getSaveFileName(self, "Export latest results table", str(self.engine.cwd / "results.csv"), "CSV (*.csv);;HTML (*.html);;LaTeX (*.tex)")
        if path:
            self.start("export results using " + self.quote(path) + ", replace")

    def open_data(self):
        if self.worker: return
        from .format_dialog import FormatDialog
        dialog = FormatDialog(self.engine.cwd, parent=self)
        if dialog.exec() == QDialog.DialogCode.Accepted and self.discard_data():
            self.start(dialog.command())

    def save_data(self):
        if self.worker:
            return
        path, _ = QFileDialog.getSaveFileName(
            self,
            "Save Stata dataset",
            str(self.engine.filename or self.engine.cwd / "data.dta"),
            "Stata data (*.dta)",
        )
        if path:
            self.start("save " + self.quote(path) + ", replace")

    def import_csv(self):
        if self.worker:
            return
        path, _ = QFileDialog.getOpenFileName(
            self, "Import CSV", str(self.engine.cwd), "CSV (*.csv);;All files (*)"
        )
        if path and self.discard_data():
            self.start("import delimited " + self.quote(path) + ", clear")

    def editor_guard(self):
        if not self.editor.document().isModified():
            return True
        answer = QMessageBox.question(
            self,
            "Unsaved do-file",
            "Save changes to this do-file?",
            QMessageBox.StandardButton.Save
            | QMessageBox.StandardButton.Discard
            | QMessageBox.StandardButton.Cancel,
            QMessageBox.StandardButton.Save,
        )
        if answer == QMessageBox.StandardButton.Cancel:
            return False
        return self.save_do() if answer == QMessageBox.StandardButton.Save else True

    def choose_working_folder(self):
        if self.worker:
            return
        folder = QFileDialog.getExistingDirectory(
            self, "Choose working folder", str(self.engine.cwd)
        )
        if folder:
            self.start("cd " + self.quote(folder))

    def use_do_folder(self):
        if self.worker:
            return
        if self.do_path is None:
            QMessageBox.information(
                self,
                "Working folder",
                "Open or save a do-file first, then choose Use do-file folder.",
            )
            return
        self.start("cd " + self.quote(str(self.do_path.resolve().parent)))

    def open_do(self):
        if not self.editor_guard():
            return
        path, _ = QFileDialog.getOpenFileName(
            self,
            "Open do-file",
            str(self.engine.cwd),
            "Stata do-files (*.do);;All files (*)",
        )
        if path:
            try:
                text = Path(path).read_text(encoding="utf-8-sig")
                self.editor.setPlainText(text)
                self.do_path = Path(path)
                self.editor.document().setModified(False)
                self.editor_title()
                self.tabs.setCurrentIndex(2)
            except (OSError, UnicodeError) as e:
                QMessageBox.warning(self, "Could not open do-file", str(e))

    def save_do(self):
        path = self.do_path
        if path is None:
            name, _ = QFileDialog.getSaveFileName(
                self,
                "Save do-file",
                str(self.engine.cwd / "analysis.do"),
                "Stata do-files (*.do)",
            )
            if not name:
                return False
            path = Path(name)
        try:
            path.write_text(self.editor.toPlainText(), encoding="utf-8")
            self.do_path = path
            self.editor.document().setModified(False)
            self.editor_title()
            return True
        except OSError as e:
            QMessageBox.warning(self, "Could not save do-file", str(e))
            return False

    def editor_title(self, *_):
        self.editor_name.setText(
            str(self.do_path or "Untitled.do")
            + (" *" if self.editor.document().isModified() else "")
        )

    def editor_location(self):
        c = self.editor.textCursor()
        self.editor.setToolTip(
            f"Line {c.blockNumber() + 1}, column {c.positionInBlock() + 1}"
        )

    def open_learning_example(self, source):
        if self.worker or not self.editor_guard():
            return
        self.do_path = None
        self.editor.setPlainText(source)
        self.editor.document().setModified(True)
        self.editor_title()
        self.tabs.setCurrentIndex(2)
        self.statusBar().showMessage(
            "Example opened. No data changed. Run its setup with Ctrl+D after saving any current dataset."
        )

    def run_editor(self):
        if self.worker:
            return
        cursor = self.editor.textCursor()
        source = (
            cursor.selectedText().replace("\u2029", "\n")
            if cursor.hasSelection()
            else self.editor.toPlainText()
        )
        if source.strip():
            self.tabs.setCurrentIndex(0)
            self.start(source)

    def save_results(self):
        path, _ = QFileDialog.getSaveFileName(
            self, "Save results", str(self.engine.cwd / "results.txt"), "Text (*.txt)"
        )
        if path:
            try:
                Path(path).write_text("".join(self.results_plain), encoding="utf-8")
            except OSError as e:
                QMessageBox.warning(self, "Could not save results", str(e))

    def export_audit(self):
        path, _ = QFileDialog.getSaveFileName(
            self,
            "Export compatibility report",
            str(self.engine.cwd / "compatibility.json"),
            "JSON (*.json)",
        )
        if path:
            try:
                Path(path).write_text(
                    json.dumps(
                        {
                            "versions": self.engine.versions,
                            "notices": self.engine.notices,
                            "commands": self.engine.ledger,
                        },
                        indent=2,
                        default=str,
                    ),
                    encoding="utf-8",
                )
            except OSError as e:
                QMessageBox.warning(self, "Could not save report", str(e))

    def closeEvent(self, event):
        if self.worker:
            QMessageBox.information(
                self,
                "Analysis is running",
                "Stop the analysis and wait for the current command to finish before closing.",
            )
            event.ignore()
            return
        if not self.editor_guard():
            event.ignore()
            return
        if self.engine.dirty:
            answer = QMessageBox.question(
                self,
                "Unsaved dataset",
                "Save the dataset before closing?",
                QMessageBox.StandardButton.Save
                | QMessageBox.StandardButton.Discard
                | QMessageBox.StandardButton.Cancel,
                QMessageBox.StandardButton.Save,
            )
            if answer == QMessageBox.StandardButton.Cancel:
                event.ignore()
                return
            if answer == QMessageBox.StandardButton.Save:
                self.save_data()
                event.ignore()
                return
        self.settings.setValue("learning_enabled", self.learning_dock.isVisible())
        self.settings.setValue("geometry", self.saveGeometry())
        self.settings.setValue("layout", self.saveState())
        self.save_working_folder()
        self.save_history()
        self.engine.close()
        event.accept()


def main(syntax=None):
    app = QApplication.instance() or QApplication(sys.argv)
    app.setApplicationName("StatBridge Qt")
    app.setOrganizationName("StatBridge")
    icon_path = Path(__file__).resolve().parent / "assets" / "statbridge.svg"
    if icon_path.is_file():
        app.setWindowIcon(QIcon(str(icon_path)))
    if sys.platform.startswith("linux") and hasattr(app, "setDesktopFileName"):
        app.setDesktopFileName("io.github.cbstatsoft.StatBridge")
    app.setStyle("Fusion")
    window = MainWindow(syntax=syntax)
    window.show()
    return app.exec()
