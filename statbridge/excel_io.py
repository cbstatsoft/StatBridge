# Copyright (C) 2026 Connor Baird
# SPDX-License-Identifier: GPL-3.0-or-later
# This program is free software under GNU GPL version 3 or later.
# It comes WITHOUT ANY WARRANTY; see LICENSE and NOTICE.
"""Excel workbook I/O with explicit dates and literal strings on export."""

import datetime as dt
import re
import numpy as np
import pandas as pd
from .syntax import StataError, options, words, unquote, qualifiers
from .result_tables import atomic_write
from .dates import EPOCH, display_date


class ExcelCommands:
    def cmd_import(self, args, opt):
        if not args.startswith("excel "):
            return super().cmd_import(args, opt)
        # firstrow is a flag for import and a valued option for export.
        o = options(opt, ("clear", "sheet", "cellrange", "firstrow", "allstring"))
        if "firstrow" in o and o["firstrow"] is not True:
            raise StataError("Import firstrow is a flag without parentheses")
        if self.dirty and "clear" not in o:
            raise StataError("Data changed; specify clear", 4)
        text = args[6:].strip()
        text = text[6:].strip() if text.startswith("using ") else text
        path = self.path(text)
        if path.suffix.lower() != ".xlsx":
            raise StataError("Excel import currently supports .xlsx, not .xls")
        from openpyxl import load_workbook
        from openpyxl.utils.cell import range_boundaries, get_column_letter

        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            sheet = str(o.get("sheet", wb.sheetnames[0]))
            if sheet not in wb.sheetnames:
                raise StataError("Worksheet not found: " + sheet)
            ws = wb[sheet]
            bounds = (1, 1, ws.max_column, ws.max_row)
            if "cellrange" in o:
                spec = str(o["cellrange"])
                if not re.fullmatch(r"[A-Z]+[1-9]\d*(?::[A-Z]+[1-9]\d*)?", spec, re.I):
                    raise StataError("cellrange() expects A1 or A1:D20")
                parsed = range_boundaries(spec.upper())
                bounds = (
                    parsed
                    if ":" in spec
                    else (parsed[0], parsed[1], ws.max_column, ws.max_row)
                )
            if bounds[0] > bounds[2] or bounds[1] > bounds[3]:
                raise StataError("Reversed Excel cell range")
            rows = list(
                ws.iter_rows(
                    min_col=bounds[0],
                    min_row=bounds[1],
                    max_col=bounds[2],
                    max_row=bounds[3],
                    values_only=True,
                )
            )
            while rows and all(x is None for x in rows[-1]):
                rows.pop()
            if not rows:
                raise StataError("Worksheet range is empty")
            names = (
                list(rows.pop(0))
                if "firstrow" in o
                else [get_column_letter(i) for i in range(bounds[0], bounds[2] + 1)]
            )
            if any(
                not isinstance(x, str) or not re.fullmatch(r"[A-Za-z_]\w{0,31}", x)
                for x in names
            ) or len(set(names)) != len(names):
                raise StataError(
                    "Headers must be unique Stata-compatible names; import without firstrow to retain column letters"
                )
            frame = pd.DataFrame(rows, columns=names)
            formats = {}
            for c in frame:
                observed = [x for x in frame[c] if pd.notna(x)]
                if "allstring" in o:
                    frame[c] = frame[c].map(lambda x: "" if pd.isna(x) else str(x))
                    continue
                if observed and all(
                    isinstance(x, (dt.datetime, dt.date, pd.Timestamp))
                    for x in observed
                ):
                    has_time = any(
                        isinstance(x, dt.datetime) and x.time() != dt.time()
                        for x in observed
                    )
                    frame[c] = frame[c].map(
                        lambda x: (
                            np.nan
                            if pd.isna(x)
                            else (
                                (
                                    pd.Timestamp(x).to_pydatetime() - EPOCH
                                ).total_seconds()
                                * 1000
                                if has_time
                                else (pd.Timestamp(x).to_pydatetime() - EPOCH).days
                            )
                        )
                    )
                    formats[c] = "%tc" if has_time else "%td"
                elif observed and all(
                    isinstance(x, (int, float, np.number, bool)) for x in observed
                ):
                    frame[c] = pd.to_numeric(frame[c]).astype(float)
                else:
                    frame[c] = frame[c].map(lambda x: "" if pd.isna(x) else str(x))
        finally:
            wb.close()
        self.df = frame
        self.labels = {}
        self.value_labels = {}
        self.label_sets = {}
        self.formats = formats
        self.file_label = ""
        self.panel = None
        self.tsvar = None
        self.filename = None
        self.changed()
        self.out(f"Imported {len(frame)} observations from {sheet}")
        self.notice(
            "EXCEL",
            "openpyxl decodes workbook date systems; date cells become Stata numeric daily or millisecond values. Mixed columns become strings. Formula cells use cached values; formulas are not evaluated. Only .xlsx; locale/type inference can differ from Stata.",
        )

    def cmd_export(self, args, opt):
        if args.startswith("results using "):
            o = options(opt, ("replace",))
            return self.export_table(args[len("results using ") :], "replace" in o)
        if not args.startswith("excel "):
            return super().cmd_export(args, opt)
        o = options(opt, ("replace", "sheet", "firstrow", "cell"))
        match = re.fullmatch(r"excel\s+(.*?)\busing\s+(.+)", args)
        if not match:
            raise StataError(
                'Syntax: export excel [varlist] using "file.xlsx", firstrow(variables) replace'
            )
        head, q = qualifiers(match[1])
        cols = self.varlist(head, default=True)
        frame = self.df.loc[self.mask(q), cols]
        if "firstrow" in o and o["firstrow"] != "variables":
            raise StataError("Only firstrow(variables) is supported")
        path = self.path(match[2])
        sheet = str(o.get("sheet", "Sheet1"))
        cell = str(o.get("cell", "A1"))
        if path.suffix.lower() != ".xlsx":
            raise StataError("Excel export requires .xlsx")
        if not re.fullmatch(r"[A-Z]+[1-9]\d*", cell, re.I):
            raise StataError("cell() expects e.g. B2")
        if not sheet or len(sheet) > 31 or re.search(r"[\\/*?:\[\]]", sheet):
            raise StataError("Invalid worksheet name")
        from openpyxl import Workbook
        from openpyxl.utils.cell import coordinate_to_tuple

        startrow, startcol = coordinate_to_tuple(cell.upper())

        def write(tmp):
            wb = Workbook()
            ws = wb.active
            ws.title = sheet
            row = startrow
            if "firstrow" in o:
                for j, c in enumerate(cols, startcol):
                    ws.cell(row, j, c)
                row += 1
            for values in frame.itertuples(index=False, name=None):
                for j, (c, value) in enumerate(zip(cols, values), startcol):
                    if pd.isna(value):
                        continue
                    target = ws.cell(row, j)
                    if self.formats.get(c) in ("%td", "%tc"):
                        value = EPOCH + dt.timedelta(
                            **(
                                {"days": float(value)}
                                if self.formats[c] == "%td"
                                else {"milliseconds": float(value)}
                            )
                        )
                        target.value = value
                        target.number_format = (
                            "yyyy-mm-dd"
                            if self.formats[c] == "%td"
                            else "yyyy-mm-dd hh:mm:ss.000"
                        )
                    else:
                        target.value = (
                            value.item() if isinstance(value, np.generic) else value
                        )
                        if isinstance(target.value, str):
                            target.data_type = "s"
                row += 1
            wb.save(tmp)
            wb.close()

        atomic_write(self, path, "replace" in o, write)
        self.out(f"Saved {path}")
        self.notice(
            "EXCEL_EXPORT",
            "Creates a workbook; replace replaces the whole file. No sheet modification/append. Strings are literal, not executable formulas. Daily/clock formats export as Excel dates; monthly/quarterly values remain numeric.",
        )

    def cmd_format(self, args, opt):
        vs, _, fmt = args.rpartition(" ")
        if fmt not in ("%td", "%tm", "%tq", "%tc"):
            return super().cmd_format(args, opt)
        options(opt, ())
        for c in self.varlist(vs, numeric=True):
            self.formats[c] = fmt
        self.dirty = True
        self.out("Date display format set; stored numeric values are unchanged")

    def cmd_list(self, args, opt):
        options(opt, ())
        head, q = qualifiers(args)
        frame = (
            self.df.loc[self.mask(q), self.varlist(head, default=True)].head(500).copy()
        )
        for c in frame:
            if self.formats.get(c) in ("%td", "%tm", "%tq", "%tc"):
                frame[c] = frame[c].map(lambda v: display_date(v, self.formats[c]))
        self.out(frame.rename(index=lambda x: x + 1).to_string())
        if self.mask(q).sum() > 500:
            self.out("Display truncated at 500 rows; use in or the Data tab.")
