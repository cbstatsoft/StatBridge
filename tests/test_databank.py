# Copyright (C) 2026 Connor Baird
# SPDX-License-Identifier: GPL-3.0-or-later
# This program is free software under GNU GPL version 3 or later.
# It comes WITHOUT ANY WARRANTY; see LICENSE and NOTICE.
from pathlib import Path
import numpy as np
import pandas as pd
import polars as pl
import pytest
from statbridge.engine import Engine
from statbridge.syntax import StataError
from statbridge.databank import convert, core, write_dataset


def source(tmp_path, name="source.csv", text=None):
    path = tmp_path / name
    path.write_text(
        text
        or "Country Name,Series Name,2020 [YR2020],2021 [YR2021]\nGhana,GDP,10,12\nGhana,Life expectancy,70,..\nKenya,GDP,20,25\nKenya,Life expectancy,65,66\n",
        encoding="utf-8",
    )
    return path


def run(engine, command):
    if not engine.execute(command):
        raise StataError(engine.events[-1]["text"], engine.scalars.get("_rc", 498))
    return engine.df


def test_worldbank_wide_preamble_footer_and_missing(tmp_path):
    p = source(tmp_path)
    p.write_text(
        "Data source,World Bank\nLast Updated,2026\n\n"
        + p.read_text()
        + "Data from database: WDI,,,,\nLast Updated: 2026,,,,\n"
    )
    # Metadata footer fields must not exceed the header's width.
    p.write_text(p.read_text().replace(",,,,", ",,,"))
    e = Engine(cwd=tmp_path)
    run(e, 'dtabnk "source.csv", clear')
    assert list(e.df) == ["Country", "Year", "GDP", "Life_expectancy"]
    assert e.df.GDP.tolist() == [10, 12, 20, 25]
    assert np.isnan(e.df.Life_expectancy.iloc[1])
    assert e.labels["Life_expectancy"] == "Life expectancy"
    assert e.model is None and e.dirty


@pytest.mark.parametrize(
    "layout,text,expected",
    [
        (
            "long",
            "Country,Year,Series,Value\nGhana,2020,GDP,12\nGhana,2021,GDP,14\n",
            [12.0, 14.0],
        ),
        (
            "wide_header_series",
            "Country Name,2020 [YR2020] - GDP [GDP.CD],2021 [YR2021] - GDP [GDP.CD]\nGhana,12,14\n",
            [12.0, 14.0],
        ),
        (
            "year_rows",
            "Year,Ghana,Côte d’Ivoire\n2020,12,15\n2021,14,18\n",
            [15.0, 18.0, 12.0, 14.0],
        ),
    ],
)
def test_other_layouts_and_original_entity_names(tmp_path, layout, text, expected):
    result = convert(source(tmp_path, text=text), {})
    assert result.layout == layout
    measure = [c for c in result.frame.columns if c not in ("Country", "Year")][0]
    assert result.frame[measure].to_list() == expected
    if layout == "year_rows":
        assert "Côte d’Ivoire" in result.frame["Country"].to_list()


def test_explicit_country_codes_and_leading_zero_ids(tmp_path):
    p = source(
        tmp_path,
        text="Country Name,Country Code,Series Code,2020,2021\nFirst,001,IND.A,1,2\nSecond,002,IND.A,3,4\n",
    )
    frame = convert(p, {"id": "Country Code", "seriescol": "Series Code"}).frame
    assert frame["Country_Code"].to_list() == ["001", "001", "002", "002"]
    assert "INDA" in frame.columns
    with pytest.raises(ValueError, match="not found"):
        convert(p, {"id": "wrong"})


def test_lazy_parquet_intermediate_cleanup_and_late_text(tmp_path, monkeypatch):
    p = source(
        tmp_path,
        text="Country,Year,Value\n"
        + "".join(f"A,{1900+i},1\n" for i in range(99))
        + "A,1999,bad\n",
    )
    paths = []

    def temporary():
        path = str(tmp_path / "temporary.parquet")
        paths.append(path)
        return path

    monkeypatch.setattr(
        core.SourceReader, "_temporary_parquet_path", staticmethod(temporary)
    )
    messages = []
    token = core.reporter.set(messages.append)
    try:
        a = convert(p, {}).frame
        b = convert(p, {"lazy": "1", "parquet": "1"}).frame
    finally:
        core.reporter.reset(token)
    assert a.equals(b) and b["Value"].null_count() == 1
    assert paths and all(not Path(x).exists() for x in paths)
    assert any("non-numeric" in m for m in messages)


def test_duplicate_policy_and_collision_safe_indicator_names(tmp_path):
    p = source(
        tmp_path,
        text="Country,Year,Series,Value\nA,2020,Year,2\nA,2020,Year,4\nA,2020,A B,8\nA,2020,A-B,9\n",
    )
    with pytest.raises(ValueError, match="duplicate|share"):
        convert(p, {})
    frame = convert(p, {"aggregate": "mean"}).frame
    assert frame["Year"].to_list() == [2020]
    assert frame["Year_1"].to_list() == [3.0]
    assert frame.width == 5
    assert frame["A_B"][0] == 8 and frame["AB"][0] == 9


@pytest.mark.parametrize(
    "join,n,countries",
    [("outer", 3, ["A", "B", "C"]), ("inner", 1, ["B"]), ("left", 2, ["A", "B"])],
)
def test_updated_upstream_merge_join_types(tmp_path, join, n, countries):
    source(tmp_path, "gdp.csv", "Country,Year,Value\nA,2020,10\nB,2020,20\n")
    source(tmp_path, "life.csv", "Country,Year,Value\nB,2020,60\nC,2020,70\n")
    e = Engine(cwd=tmp_path)
    run(e, f'dtabnk "gdp.csv" "life.csv", merge join({join}) clear')
    assert len(e.df) == n and sorted(e.df.Country) == countries
    assert set(e.df) == {"Country", "Year", "gdp", "life"}
    if join == "outer":
        assert np.isnan(e.df.set_index("Country").loc["A", "life"])
        assert np.isnan(e.df.set_index("Country").loc["C", "gdp"])


def test_merge_name_conflicts_suffix_or_error(tmp_path):
    for name in ("one.csv", "two.csv"):
        source(tmp_path, name, "Country,Year,Series,Value\nA,2020,GDP,2\n")
    e = Engine(cwd=tmp_path)
    run(e, 'dtabnk "one.csv" "two.csv", merge clear')
    assert set(e.df) == {"Country", "Year", "GDP", "GDP_two"}
    old = e.df.copy()
    with pytest.raises(StataError, match="conflict"):
        run(e, 'dtabnk "one.csv" "two.csv", merge clear conflicts(error)')
    pd.testing.assert_frame_equal(old, e.df)


def test_long_output_merge_requires_extra_series_key(tmp_path):
    source(tmp_path, "a.csv")
    source(tmp_path, "b.csv")
    e = Engine(cwd=tmp_path)
    with pytest.raises(StataError, match="duplicate merge keys"):
        run(e, 'dtabnk "a.csv" "b.csv", merge long clear')
    run(e, 'dtabnk "a.csv" "b.csv", merge long on(Country Year Series) clear')
    assert len(e.df) == 8


@pytest.mark.parametrize("bad", ["A,2020Q1,1", ",2020,1", "A,2020.5,1"])
def test_invalid_keys_rejected_without_replacing_data(tmp_path, bad):
    source(tmp_path, text="Country,Year,Value\n" + bad + "\n")
    e = Engine(cwd=tmp_path)
    run(e, "set obs 3")
    run(e, "generate original = _n")
    old = e.df.copy()
    revision = e.revision
    with pytest.raises(StataError):
        run(e, 'dtabnk "source.csv", clear')
    pd.testing.assert_frame_equal(e.df, old)
    assert e.revision == revision


def test_preview_and_conversion_only_preserve_model_and_dataset(tmp_path):
    source(tmp_path)
    e = Engine(cwd=tmp_path)
    e.run("set obs 20\ngenerate x = _n\ngenerate y = 2*x+sin(x)\nregress y x")
    old = e.df.copy()
    model = e.model
    revision = e.revision
    with pytest.raises(StataError, match="Data changed"):
        run(e, 'dtabnk "source.csv"')
    run(e, 'dtabnk "source.csv", preview rows(2)')
    assert len(e.results_table) == 2 and e.r["N"] == 4
    run(e, 'dtabnk "source.csv", saving("out.parquet") noload')
    pd.testing.assert_frame_equal(e.df, old)
    assert e.model is model and e.revision == revision
    assert pl.read_parquet(tmp_path / "out.parquet").height == 4
    with pytest.raises(StataError):
        run(e, 'dtabnk "source.csv", preview saving("never.dta")')
    assert not (tmp_path / "never.dta").exists()


@pytest.mark.parametrize("suffix", ["dta", "sav", "parquet", "RData"])
def test_export_formats_roundtrip_and_overwrite_guard(tmp_path, suffix):
    if suffix == "RData":
        pytest.importorskip("pyreadr")
    source(tmp_path)
    e = Engine(cwd=tmp_path)
    run(e, f'dtabnk "source.csv", clear saving("panel.{suffix}")')
    path = tmp_path / f"panel.{suffix}"
    before = path.read_bytes()
    with pytest.raises(StataError):
        run(e, f'dtabnk "source.csv", clear saving("panel.{suffix}")')
    assert path.read_bytes() == before
    if suffix == "parquet":
        frame = pl.read_parquet(path).to_pandas()
    elif suffix == "RData":
        import pyreadr

        frame = pyreadr.read_r(str(path))["df"].reset_index(drop=True)
    else:
        import pyreadstat

        frame, _ = (
            pyreadstat.read_dta(str(path))
            if suffix == "dta"
            else pyreadstat.read_sav(str(path))
        )
    assert frame.GDP.tolist() == [10.0, 12.0, 20.0, 25.0]
    assert np.isnan(frame.Life_expectancy.iloc[1])


@pytest.mark.parametrize("suffix", ["xlsx", "xls"])
def test_excel_sheet_header_and_allstring_source(tmp_path, suffix):
    path = tmp_path / f"book.{suffix}"
    rows = [
        ["note", "preamble"],
        ["Country", "Year", "Value"],
        ["001", 2020, 12],
        ["002", 2020, ".."],
    ]
    if suffix == "xlsx":
        from openpyxl import Workbook

        wb = Workbook()
        wb.active.title = "Unused"
        ws = wb.create_sheet("Panel")
        for row in rows:
            ws.append(row)
        wb.save(path)
    else:
        xlwt = pytest.importorskip("xlwt")
        wb = xlwt.Workbook()
        wb.add_sheet("Unused")
        ws = wb.add_sheet("Panel")
        for i, row in enumerate(rows):
            for j, value in enumerate(row):
                ws.write(i, j, value)
        wb.save(str(path))
    result = convert(path, {"sheet": "Panel"}).frame
    assert result["Country"].to_list() == ["001", "002"]
    assert result["Value"].to_list() == [12.0, None]


def test_memory_refusal_preserves_state(tmp_path, monkeypatch):
    source(tmp_path)
    e = Engine(cwd=tmp_path)
    run(e, "set obs 2")
    old = e.df.copy()
    monkeypatch.setattr(
        core.MemoryManager, "available_ram_mb", staticmethod(lambda: 20)
    )
    with pytest.raises(StataError, match="RAM|Refusing"):
        run(e, 'dtabnk "source.csv", clear safe')
    pd.testing.assert_frame_equal(e.df, old)


def test_atomic_write_failure_keeps_existing_file(tmp_path, monkeypatch):
    import pyreadstat

    e = Engine()
    path = tmp_path / "existing.dta"
    path.write_bytes(b"original")

    def fail(frame, name, **kwargs):
        Path(name).write_bytes(b"partial")
        raise RuntimeError("disk failure")

    monkeypatch.setattr(pyreadstat, "write_dta", fail)
    with pytest.raises(RuntimeError, match="disk failure"):
        write_dataset(e, pl.DataFrame({"x": [1]}), path, replace=True)
    assert path.read_bytes() == b"original"
    assert list(tmp_path.iterdir()) == [path]


def test_completion_paths_and_options(tmp_path):
    from statbridge.completion import Suggestions
    from statbridge.path_completion import path_completions

    source(tmp_path, "panel.csv")
    (tmp_path / "wrong.dta").touch()
    for prefix in ('dtabnk "pa', 'dtabnk "first.csv" "pa'):
        values = path_completions(prefix, len(prefix), tmp_path)
        assert values and any("panel.csv" in x[0] for x in values)
    suggestions = Suggestions()
    assert 'dtabnk "x", join(inner)' in suggestions.candidates('dtabnk "x", join(in')
    assert 'dtabnk "x", layout(wide)' in suggestions.candidates('dtabnk "x", layout(wi')


def test_dialog_preview_runs_off_thread_and_generates_merge_command(tmp_path):
    import os, time

    os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
    from PySide6.QtWidgets import QApplication
    from statbridge.databank_ui import DatabankDialog

    app = QApplication.instance() or QApplication([])
    a = source(tmp_path, "a.csv")
    b = source(tmp_path, "b.csv")
    dialog = DatabankDialog([a, b], tmp_path)
    assert "merge on(Country Year)" in dialog.command()
    dialog.preview()
    deadline = time.monotonic() + 15
    while dialog.worker is not None and time.monotonic() < deadline:
        app.processEvents()
        time.sleep(0.01)
    assert dialog.worker is None
    assert dialog.grid.rowCount() == 4
    assert "4 observations" in dialog.status.toPlainText()
    dialog.close()


def test_reserved_indicator_names_do_not_overwrite_entity_or_keys(tmp_path):
    p = source(
        tmp_path,
        text="Country Name,Series Name,2020\nGhana,Country,3\nGhana,Series,4\nGhana,Series_Code,5\n",
    )
    frame = convert(p, {}).frame
    assert frame["Country"].to_list() == ["Ghana"]
    assert frame["Country_1"].to_list() == [3.0]
    assert frame["Series"].to_list() == [4.0]
    assert frame["Series_Code"].to_list() == [5.0]


def test_raw_header_mapping_disambiguates_sanitised_collisions(tmp_path):
    p = source(tmp_path, text="Country,A-B,AB,Value\nGhana,ignore,2020,7\n")
    frame = convert(p, {"layout": "long", "yearcol": "AB"}).frame
    assert frame["Year"].to_list() == [2020]
    assert frame["Value"].to_list() == [7.0]


def test_missing_year_refused_instead_of_dropping_observation(tmp_path):
    p = source(tmp_path, text="Country,Year,Value\nGhana,,7\n")
    with pytest.raises(ValueError, match="invalid or non-annual"):
        convert(p, {})


def test_merge_labels_follow_renamed_indicator_columns(tmp_path):
    source(tmp_path, "one.csv", "Country,Year,Series,Value\nGhana,2020,A-B,1\n")
    source(tmp_path, "two.csv", "Country,Year,Series,Value\nGhana,2020,AB,2\n")
    e = Engine(cwd=tmp_path)
    run(e, 'dtabnk "one.csv" "two.csv", merge clear')
    assert e.labels["AB"] == "A-B"
    assert e.labels["AB_two"] == "AB"
