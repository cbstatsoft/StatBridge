# Copyright (C) 2026 Connor Baird
# SPDX-License-Identifier: GPL-3.0-or-later
# This program is free software under GNU GPL version 3 or later.
# It comes WITHOUT ANY WARRANTY; see LICENSE and NOTICE.
import copy
import numpy as np
import pandas as pd
import pytest
import statsmodels.api as sm
from scipy import special, stats
from statbridge.engine import Engine


def ok(e, cmd):
    assert e.run(cmd), [x["text"] for x in e.events if x["kind"] == "error"][-2:]


def sample(tmp_path):
    rng = np.random.default_rng(52)
    n = 360
    e = Engine(cwd=tmp_path)
    e.df = pd.DataFrame(
        {
            "x": rng.normal(size=n),
            "z": rng.normal(size=n),
            "g": np.tile([0.0, 1.0, 2.0], 120),
        }
    )
    e.df["y"] = (
        1
        + 2 * e.df.x
        + e.df.z
        + 1.5 * (e.df.g == 1)
        + 0.8 * e.df.x * (e.df.g == 1)
        + rng.normal(size=n)
    )
    return e


def test_margins_interaction_prediction_and_analytic_covariance(tmp_path):
    e = sample(tmp_path)
    ok(e, "regress y c.x##i.g z, vce(robust)")
    m = e.model
    before = copy.deepcopy(e.e)
    ok(e, "margins g, at(x=(0 2))")
    table = e.margins_result["table"]
    for _, row in table.iterrows():
        data = e.df.copy()
        data["x"] = row.x
        data["g"] = row.g
        design = pd.DataFrame(
            {
                "_cons": 1.0,
                "x": data.x,
                "1.g": (data.g == 1).astype(float),
                "2.g": (data.g == 2).astype(float),
                "x#1.g": data.x * (data.g == 1),
                "x#2.g": data.x * (data.g == 2),
                "z": data.z,
            }
        )[m.params.index]
        avg = design.mean().to_numpy()
        assert row.Estimate == pytest.approx(avg @ m.params)
        assert row.SE**2 == pytest.approx(avg @ m.cov @ avg)
    assert e.model is m and e.e == before
    ok(e, "margins, dydx(x) over(g)")
    assert np.allclose(
        e.results_table.Estimate,
        [m.params.x, m.params.x + m.params["x#1.g"], m.params.x + m.params["x#2.g"]],
    )
    assert not e.run("margins, at(g=8)")


@pytest.mark.parametrize("command", ["logit", "probit"])
def test_nonlinear_margins_delta_and_factor_contrast(command, tmp_path):
    e = sample(tmp_path)
    rng = np.random.default_rng(31)
    e.df["binary"] = (
        rng.uniform(size=len(e.df))
        < special.expit(0.1 + 0.6 * e.df.x + 0.4 * (e.df.g == 1))
    ).astype(float)
    ok(e, f"{command} binary c.x##ib1.g z, vce(robust)")
    m = e.model
    ok(e, "margins, dydx(x g) at(x=1)")
    link = special.expit if command == "logit" else stats.norm.cdf

    def f(beta):
        data = e.df.copy()
        data.x = 1
        # Explicit frozen base 1; reconstruct independently.
        x = pd.DataFrame(
            {
                "_cons": 1.0,
                "x": data.x,
                "0.g": (data.g == 0).astype(float),
                "2.g": (data.g == 2).astype(float),
                "x#0.g": data.x * (data.g == 0),
                "x#2.g": data.x * (data.g == 2),
                "z": data.z,
            }
        )[m.params.index].to_numpy()
        b = np.asarray(beta)
        eta = x @ b
        slope = (
            b[list(m.params.index).index("x")]
            + (data.g == 0).to_numpy() * b[list(m.params.index).index("x#0.g")]
            + (data.g == 2).to_numpy() * b[list(m.params.index).index("x#2.g")]
        )
        derivative = (
            link(eta) * (1 - link(eta)) if command == "logit" else stats.norm.pdf(eta)
        )
        return np.mean(derivative * slope)

    row = e.results_table.iloc[0]
    b = m.params.to_numpy()
    eps = 1e-5
    grad = np.array(
        [
            (f(b + eps * np.eye(len(b))[i]) - f(b - eps * np.eye(len(b))[i]))
            / (2 * eps)
            for i in range(len(b))
        ]
    )
    assert row.Estimate == pytest.approx(f(b))
    assert row.SE**2 == pytest.approx(grad @ m.cov @ grad, rel=1e-7)
    assert e.results_table.Effect.tolist() == ["dy/dx", "0.g vs 1", "2.g vs 1"]
    assert not e.run("margins, at(g=99)")


def panel(tmp_path):
    rng = np.random.default_rng(10)
    e = Engine(cwd=tmp_path)
    n = 80
    t = 5
    e.df = pd.DataFrame(
        {"id": np.repeat(np.arange(n), t), "time": np.tile(np.arange(t), n)}
    )
    effect = np.repeat(rng.normal(size=n), t)
    e.df["x"] = 0.7 * effect + rng.normal(size=n * t)
    e.df["y"] = 1 + 2 * e.df.x + effect + 0.2 * e.df.time + rng.normal(size=n * t)
    ok(e, "xtset id time")
    return e


def test_mundlak_balanced_slope_equals_fe_and_preserves_active_model(tmp_path):
    e = panel(tmp_path)
    ok(e, "xtreg y x i.time, fe vce(cluster id)")
    fe = e.model
    e_state = copy.deepcopy(e.e)
    ok(e, "estat mundlak")
    assert e.model is fe and e.e == e_state and e.r["df"] == 1 and e.r["p"] < 0.05
    ok(e, "xtreg y x i.time, cre vce(cluster id)")
    assert e.model.params.x == pytest.approx(fe.params.x, rel=1e-9)
    ok(e, "estat mundlak")
    names = e.model.metadata["cre_means"]
    b = e.model.params[names]
    v = e.model.cov.loc[names, names]
    assert e.r["chi2"] == pytest.approx(b @ np.linalg.solve(v, b))
    ok(e, "xtreg y x if _n > 1, re")
    m = e.model
    assert not e.run("estat mundlak") and e.model is m
    assert not e.run("xtreg y x, cre fe")


@pytest.mark.parametrize("vce", ["", " , vce(robust)", " , vce(cluster g)"])
def test_iv_diagnostics_distributions_and_unsupported_cases(vce, tmp_path):
    e = sample(tmp_path)
    rng = np.random.default_rng(112)
    e.df["iv1"] = rng.normal(size=len(e.df))
    e.df["iv2"] = rng.normal(size=len(e.df))
    e.df.x = e.df.iv1 + 0.6 * e.df.iv2 + 0.4 * e.df.z + rng.normal(size=len(e.df))
    e.df.y = 1 + 2 * e.df.x + e.df.z + rng.normal(size=len(e.df))
    ok(e, "ivregress 2sls y z (x = iv1 iv2)" + vce)
    m = e.model
    ok(e, "estat firststage")
    table = e.results_table
    assert table.loc["x", "f.dist"].startswith("F" if not vce else "chi2")
    ok(e, "estat endogenous")
    assert e.r["distribution"].startswith("F" if not vce else "chi2")
    assert e.model is m
    if "cluster" in vce:
        assert not e.run("estat overid")
    else:
        ok(e, "estat overid")
        assert e.r["df"] == 1
    ok(e, "ivregress 2sls y z (x = iv1)" + vce)
    assert not e.run("estat overid")


def test_absorbed_slopes_match_dummy_ols_and_two_way_clusters(tmp_path):
    e = panel(tmp_path)
    dummies = pd.get_dummies(
        e.df[["id", "time"]].astype("category"), drop_first=True, dtype=float
    )
    full = sm.add_constant(pd.concat([e.df[["x"]], dummies], axis=1))
    ordinary = sm.OLS(e.df.y, full).fit()
    ok(e, "reghdfe y x, absorb(id time)")
    assert e.model.params.x == pytest.approx(ordinary.params.x, rel=1e-9)
    assert np.allclose(e.model.residual, ordinary.resid, atol=1e-9)
    ok(e, "reghdfe y x, absorb(id time) vce(cluster id time)")
    assert e.model.df == 4 and np.isfinite(e.model.cov).all().all()
    assert e.model.metadata["singletons_dropped"] == 0
    assert not e.run("reghdfe y x, absorb(id time) vce(cluster id time g)")


def test_hdfe_iterative_singleton_removal(tmp_path):
    e = panel(tmp_path)
    extra = pd.DataFrame(
        {
            "id": [80, 81, 81],
            "time": [0, 10, 11],
            "x": [1.0, 2.0, 3.0],
            "y": [5.0, 6.0, 7.0],
        }
    )
    e.df = pd.concat([e.df, extra], ignore_index=True)
    ok(e, "reghdfe y x, absorb(id time)")
    assert len(e.model.sample) == 400 and e.model.metadata["singletons_dropped"] == 3


def test_two_period_did_matches_difference_in_means_and_refuses_staggered(tmp_path):
    e = panel(tmp_path)
    e.df = e.df[e.df.time.isin([0, 1])].reset_index(drop=True)
    e.df["treated"] = ((e.df.id >= 40) & (e.df.time == 1)).astype(float)
    e.df["y"] = (
        e.df.id / 10 + 3 * e.df.time + 5 * e.df.treated + np.sin(e.df.id) * e.df.time
    )
    ok(e, "xtset id time")
    means = e.df.groupby([e.df.id >= 40, "time"]).y.mean()
    expected = (means[True, 1] - means[True, 0]) - (means[False, 1] - means[False, 0])
    for cmd in ("didregress", "xtdidregress"):
        ok(e, f"{cmd} (y) (treated), group(id) time(time)")
        assert e.model.params.treated == pytest.approx(expected)
        assert e.model.command == cmd
    e.df.loc[0, "treated"] = 1
    assert not e.run("didregress (y) (treated), group(id) time(time)")


def test_tables_atomic_exports_and_literal_html(tmp_path):
    e = sample(tmp_path)
    ok(
        e,
        "regress y x\nestimates store first\nregress y x z, vce(robust)\nestimates store second",
    )
    ok(e, 'etable, estimates(first second) export("models.html")')
    assert e.results_table.loc["x", "first"] == e.stored["first"].params.x
    original = (tmp_path / "models.html").read_bytes()
    assert not e.run('etable, export("models.html")')
    assert (tmp_path / "models.html").read_bytes() == original
    ok(e, 'esttab first second using "models.csv", se')
    exported = pd.read_csv(tmp_path / "models.csv", index_col=0)
    assert exported.loc["x", "second"] == pytest.approx(e.stored["second"].params.x)
    ok(e, 'dtable x y, by(g) export("summary.tex")')
    assert e.results_table.loc[(0.0, "x"), "N"] == 120
    e.df["label"] = "<script>alert(1)</script>"
    ok(e, 'dtable y, by(label) export("escaped.html")')
    assert "<script>" not in (tmp_path / "escaped.html").read_text()


def test_dates_masks_negative_epoch_and_display(tmp_path):
    from statbridge.dates import display_date

    e = Engine(cwd=tmp_path)
    e.df = pd.DataFrame(
        {"text": ["1960-01-01", "1959-12-31", "2020-02-29", "2021-02-29"]}
    )
    ok(e, 'generate d = daily(text,"YMD")\nformat d %td')
    assert e.df.d.iloc[:2].tolist() == [0, -1] and np.isnan(e.df.d.iloc[3])
    ok(e, "generate month = mofd(d)\ngenerate back = dofm(month)")
    assert e.df.month.iloc[1] == -1 and e.df.back.iloc[1] == -31
    ok(
        e,
        'generate q = quarterly("1959q4","YQ")\ngenerate clocktime = clock("1960-01-01 00:00:01.500","YMDhms")',
    )
    assert e.df.q.eq(-1).all() and e.df.clocktime.eq(1500).all()
    assert (
        display_date(-1, "%td") == "31dec1959" and display_date(-1, "%tm") == "1959m12"
    )
    assert display_date("hello", None) is None
    assert not e.run('generate bad = daily(text,"YD")')
    ok(e, 'save "dates.dta"\nuse "dates.dta", clear')
    assert e.formats["d"] == "%td" and e.df.d.iloc[1] == -1


def test_excel_dates_ranges_literal_formulas_and_guards(tmp_path):
    from openpyxl import Workbook, load_workbook
    from openpyxl.utils.datetime import CALENDAR_MAC_1904
    import datetime as dt

    wb = Workbook()
    wb.epoch = CALENDAR_MAC_1904
    ws = wb.active
    ws.title = "Source"
    ws.append(["x", "when", "mixed"])
    ws.append([2, dt.datetime(2020, 2, 29), "=1+1"])
    ws.append([3, dt.datetime(2020, 3, 1), "word"])
    wb.save(tmp_path / "source.xlsx")
    e = Engine(cwd=tmp_path)
    ok(e, 'import excel "source.xlsx", sheet(Source) cellrange(A1:C3) firstrow clear')
    assert (
        e.formats["when"] == "%td"
        and e.df.x.tolist() == [2, 3]
        and e.df.mixed.iloc[0] == ""
    )
    e.df.loc[0, "mixed"] = "=SUM(1,2)"
    ok(e, 'export excel using "copy.xlsx", firstrow(variables) sheet(Results) cell(B2)')
    check = load_workbook(tmp_path / "copy.xlsx", data_only=False)
    assert (
        check["Results"]["D3"].value == "=SUM(1,2)"
        and check["Results"]["D3"].data_type == "s"
    )
    assert check["Results"]["C3"].value == dt.datetime(2020, 2, 29)
    assert not e.run('export excel using "copy.xlsx"')
    ok(e, 'import excel "copy.xlsx", sheet(Results) cellrange(B2:D4) firstrow clear')
    assert e.df.mixed.iloc[0] == "=SUM(1,2)"


def test_graphs_and_last_successful_margins(tmp_path):
    e = sample(tmp_path)
    ok(e, "regress y x\nmargins, at(x=(-1(1)1))\nmarginsplot")
    assert len(e.figure.axes[0].lines) > 0
    for cmd in (
        'twoway (scatter y x) (lfit y x), title("Fit")',
        "rvfplot",
        "qnorm x",
        'twoway scatter y x || line y x, title("Layers")',
    ):
        ok(e, cmd)
        assert e.figure.axes
    ok(e, 'graph export "graph.svg"')
    assert (tmp_path / "graph.svg").stat().st_size > 1000
    ok(e, "clear all")
    assert not e.run("marginsplot")


def test_multiway_covariance_matches_cluster_score_formula(tmp_path):
    e = panel(tmp_path)
    ok(e, "reghdfe y x, absorb(id time) vce(cluster id time)")
    m = e.model
    n = len(e.df)
    # Balanced two-way demeaning, independent of the absorber backend.
    x = (
        e.df.x
        - e.df.groupby("id").x.transform("mean")
        - e.df.groupby("time").x.transform("mean")
        + e.df.x.mean()
    )
    score = x * m.residual

    def meat(labels):
        sums = score.groupby(labels).sum()
        g = len(sums)
        return float((sums**2).sum()) * g / (g - 1) * (n - 1) / (n - 1)

    joint = pd.Series(list(zip(e.df.id, e.df.time)))
    variance = (meat(e.df.id) + meat(e.df.time) - meat(joint)) / float(x @ x) ** 2
    assert m.cov.loc["x", "x"] == pytest.approx(variance, rel=1e-8)


def test_new_completion_and_excel_paths(tmp_path):
    from statbridge.completion import Suggestions

    s = Suggestions()
    s.variables = ["x", "group", "time"]
    s.numeric_variables = ["x", "group", "time"]
    s.estimates = ["first", "second"]
    s.cwd = tmp_path
    assert "estat mundlak" in s.candidates("estat mund")
    assert "margins, dydx(group" in s.candidates("margins, dydx(gr")
    assert "etable, estimates(first" in s.candidates("etable, estimates(fir")
    assert "reghdfe y x, absorb(group" in s.candidates("reghdfe y x, absorb(gro")
    assert "import excel" in s.candidates("import exc")
    assert not any(
        "delimiter(" in x for x in s.candidates('import excel "book.xlsx", del')
    )
    (tmp_path / "book.xlsx").write_bytes(b"")
    assert any("book.xlsx" in x for x in s.candidates('import excel "boo'))


def test_failed_import_leaves_current_data_and_failed_absorb_fit_keeps_model(tmp_path):
    e = sample(tmp_path)
    ok(e, "regress y x")
    model = e.model
    data = e.df.copy()
    assert not e.run('import excel "absent.xlsx", firstrow clear')
    pd.testing.assert_frame_equal(e.df, data)
    assert e.model is model
    assert not e.run("reghdfe y x, absorb(g) vce(cluster missing_cluster)")
    assert e.model is model
    assert not e.run("margins, at(x=(0)) unknown")
